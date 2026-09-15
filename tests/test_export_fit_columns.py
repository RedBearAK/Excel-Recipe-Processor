"""
fit_columns on export: exact agreement with format_excel's auto_fit,
the single-sheet export bridge, and targeted null clearing.

File: excel_recipe_processor/tests/test_export_fit_columns.py

Run: PYTHONPATH=. python3 tests/test_export_fit_columns.py
"""

import os
import sys
import shutil
import tempfile

import pandas as pd
from openpyxl import load_workbook

from excel_recipe_processor.core.recipe_pipeline import RecipePipeline
from excel_recipe_processor.core.workbook_session import WorkbookSession
from excel_recipe_processor.writers.excel_writer import clear_null_cells, fitted_widths


WORK = tempfile.mkdtemp(prefix='erp_fit_')


def frame() -> pd.DataFrame:
    return pd.DataFrame({
        'Id': ['P00000001', 'P00000002', 'P00000003', None],
        'A very long header name indeed': ['x', 'y', 'z', 'w'],
        'Weight': [1285.0, 53.1, 0.0, 12345.678901],
        'Stamp': pd.to_datetime(['2025-08-04 05:32:00', '2025-08-05 16:22:00', None, '2025-08-06 00:00:00']),
        'Count': pd.array([1, 2, None, 4], dtype='Int64'),
        'Note': ['', None, 'CV/EYE', 'a note that is quite a bit longer than the others in this column'],
    })


def recipe(fit: bool, autofit_after: bool, target: str) -> str:
    csv_path = os.path.join(WORK, 'frame.parquet')
    frame().to_parquet(csv_path)
    fit_keys = '''    fit_columns: true
    fit_for_auto_filter: true
''' if fit else ''
    formatting = '''
  - step_description: "format"
    processor_type: "format_excel"
    target_file: "__T__"
    formatting:
      - sheet_names: ["Data"]
        header_bold: true
        auto_filter: true
''' + ('        auto_fit_columns: true\n        min_column_width: 8\n        max_column_width: 100\n' if autofit_after else '')
    return f'''
settings:
  description: "fit test"
  stages:
    - stage_name: "stg_f"
      description: "f"
      protected: false
recipe:
  - step_description: "in"
    processor_type: "import_file"
    input_file: "{csv_path}"
    save_to_stage: "stg_f"
  - step_description: "out"
    processor_type: "export_file"
    source_stage: "stg_f"
    output_file: "__T__"
    sheet_name: "Data"
    create_backup: false
{fit_keys}{formatting}'''.replace('__T__', target)


def run(text: str, name: str):
    path = os.path.join(WORK, name)
    with open(path, 'w', encoding='utf-8') as handle:
        handle.write(text)
    WorkbookSession.reset()
    RecipePipeline().run_complete_recipe(path)


def widths_of(path: str) -> dict:
    sheet = load_workbook(path)['Data']
    return {letter: dimension.width for letter, dimension in sheet.column_dimensions.items()}


def test_fitted_widths_equal_auto_fit() -> bool:
    print('\nWidths from fit_columns on export equal what format_excel auto_fit computes on the same workbook...')
    fast = os.path.join(WORK, 'fast.xlsx')
    slow = os.path.join(WORK, 'slow.xlsx')
    run(recipe(fit=True, autofit_after=False, target=fast), 'fast.yaml')
    run(recipe(fit=False, autofit_after=True, target=slow), 'slow.yaml')
    a, b = widths_of(fast), widths_of(slow)
    letters = sorted(set(a) | set(b))
    diffs = [(letter, a.get(letter), b.get(letter)) for letter in letters if a.get(letter) != b.get(letter)]
    for letter, fast_width, slow_width in diffs:
        print(f"  {letter}: fit_columns={fast_width} auto_fit={slow_width}")
    good = not diffs and len(a) == 6
    print(f"  {len(a)} columns, disagreements={len(diffs)} -> {'OK' if good else 'FAIL'}")
    return good


def test_single_sheet_export_stays_in_the_session() -> bool:
    print('\nA single-sheet export under a session is bridged: no file on disk until run end, one save...')
    target = os.path.join(WORK, 'bridged.xlsx')
    text = recipe(fit=True, autofit_after=False, target=target)
    # a debug step after the export would see the session copy; here we check
    # the file does NOT exist at the moment the format step starts by making
    # the format step the observer: it logs "Loading Excel file" only when it
    # has to read from disk
    import logging, io

    class Catch(logging.Handler):
        def __init__(self):
            super().__init__()
            self.lines = []

        def emit(self, record):
            self.lines.append(record.getMessage())

    catcher = Catch()
    logging.getLogger().addHandler(catcher)
    try:
        run(text, 'bridged.yaml')
    finally:
        logging.getLogger().removeHandler(catcher)
    reloaded = any('Loading Excel file' in line for line in catcher.lines)
    exists = os.path.isfile(target)
    good = exists and not reloaded
    print(f"  file written at run end={exists} format step reloaded from disk={reloaded} -> {'OK' if good else 'FAIL'}")
    return good


def test_null_clearing_visits_only_nulls() -> bool:
    print('\nclear_null_cells touches exactly the cells the frame says were null...')
    from openpyxl import Workbook
    df = frame()
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(df.columns))
    for row in df.itertuples(index=False):
        sheet.append(['' if pd.isna(value) else value for value in row])
    cleared = clear_null_cells(sheet, df)
    expected = int(df.isna().to_numpy().sum())
    id_cell = sheet.cell(row=5, column=1).value
    note_empty_string_kept = sheet.cell(row=2, column=6).value == ''      # '' in the frame is NOT null
    good = cleared == expected and id_cell is None and note_empty_string_kept
    print(f"  cleared={cleared} expected={expected} null Id now None={id_cell is None} real '' kept={note_empty_string_kept} -> {'OK' if good else 'FAIL'}")
    return good


def main() -> int:
    tests = [test_fitted_widths_equal_auto_fit, test_single_sheet_export_stays_in_the_session, test_null_clearing_visits_only_nulls]
    passed = 0
    try:
        for test in tests:
            if test():
                passed += 1
    finally:
        shutil.rmtree(WORK, ignore_errors=True)
    print(f"\n{passed}/{len(tests)} tests passed")
    return 0 if passed == len(tests) else 1


if __name__ == '__main__':
    sys.exit(main())


# End of file #
