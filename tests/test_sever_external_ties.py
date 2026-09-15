"""
Antagonistic tests for sever_external_ties.

tests/test_sever_external_ties.py

Three hand-built packages drive the surgery: the audit drill's hostile
file (refs a sheet that is NOT local and has no cached value -> must
refuse, all-or-nothing), a copied-tab file where every [N] names a
local sheet (-> everything retargets, plumbing gone, Excel-loadable),
and a freeze file where nothing retargets (-> formulas frozen with
shared slaves, names / CF / DV dropped, DV count fixed). Runnable
directly or with pytest.
"""

import os
import re
import sys
import json
import hashlib
import tempfile
import zipfile

import openpyxl

from excel_recipe_processor.core.base_processor import StepProcessorError
from excel_recipe_processor.processors.sever_external_ties_processor import (
    SeverExternalTiesProcessor,
)
from excel_recipe_processor.processors._helpers.external_ties_inventory import (
    inventory_external_ties,
)

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from test_audit_external_ties import (  # noqa: E402
    NS_R,
    REL,
    NS_PKG,
    NS_MAIN,
    build_hostile_workbook,
)


def build_linked_workbook(path: str, local_sheet_for_refs: bool) -> None:
    """
    Two sheets, Data and Ext, plus link [1] whose refs name either
    'Ext' (retargetable) or 'Gone' (must freeze/drop). Every carrier
    kind present, every formula with a cached value; F1:F3 shared.
    """
    target = 'Ext' if local_sheet_for_refs else 'Gone'
    xml_head = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    xml_workbook = (
        f'{xml_head}<workbook xmlns="{NS_MAIN}" xmlns:r="{NS_R}">'
        '<sheets><sheet name="Data" sheetId="1" r:id="rId1"/>'
        '<sheet name="Ext" sheetId="2" r:id="rId2"/></sheets>'
        '<externalReferences><externalReference r:id="rId3"/></externalReferences>'
        '<definedNames>'
        '<definedName name="LocalName">Data!$A$1</definedName>'
        f'<definedName name="ExtName" hidden="1">[1]{target}!$B$2</definedName>'
        '</definedNames>'
        '<calcPr calcId="1" calcMode="manual"/>'
        '</workbook>'
    )
    xml_workbook_rels = (
        f'{xml_head}<Relationships xmlns="{NS_PKG}">'
        f'<Relationship Id="rId1" Type="{REL}worksheet" Target="worksheets/sheet1.xml"/>'
        f'<Relationship Id="rId2" Type="{REL}worksheet" Target="worksheets/sheet2.xml"/>'
        f'<Relationship Target="externalLinks/externalLink1.xml" Type="{REL}externalLink" Id="rId3"/>'
        f'<Relationship Id="rId4" Type="{REL}calcChain" Target="calcChain.xml"/>'
        '</Relationships>'
    )
    xml_data = (
        f'{xml_head}<worksheet xmlns="{NS_MAIN}"><sheetData>'
        '<row r="1">'
        '<c r="A1"><v>3</v></c>'
        f'<c r="B1"><f>[1]{target}!A1*2</f><v>14</v></c>'
        f"<c r=\"C1\" t=\"str\"><f>'[1]{target}'!A1&amp;\"x\"</f><v>7x</v></c>"
        '<c r="D1"><f>[1]!ExtOnlyName</f><v>5</v></c>'
        '<c r="E1"><f>[1]!LocalName</f><v>3</v></c>'
        f'<c r="F1"><f t="shared" ref="F1:F3" si="0">[1]{target}!A1+A1</f><v>10</v></c>'
        '<c r="G1"><f>A1*10</f><v>30</v></c>'
        '</row>'
        '<row r="2"><c r="F2"><f t="shared" si="0"/><v>11</v></c></row>'
        '<row r="3"><c r="F3"><f t="shared" si="0"/><v>12</v></c></row>'
        '</sheetData>'
        '<conditionalFormatting sqref="A1:A9">'
        f'<cfRule type="expression" priority="1"><formula>A1&gt;[1]{target}!A1</formula></cfRule>'
        '</conditionalFormatting>'
        '<conditionalFormatting sqref="G1:G9">'
        '<cfRule type="expression" priority="2"><formula>G1&gt;0</formula></cfRule>'
        '</conditionalFormatting>'
        '<dataValidations count="2">'
        f'<dataValidation type="list" sqref="B2:B9"><formula1>[1]{target}!$A$1:$A$5</formula1></dataValidation>'
        '<dataValidation type="list" sqref="C2:C9"><formula1>"a,b"</formula1></dataValidation>'
        '</dataValidations>'
        '</worksheet>'
    )
    xml_ext = (
        f'{xml_head}<worksheet xmlns="{NS_MAIN}"><sheetData>'
        '<row r="1"><c r="A1"><v>7</v></c></row>'
        '<row r="2"><c r="B2"><v>8</v></c></row>'
        '</sheetData></worksheet>'
    )
    xml_link = (
        f'{xml_head}<externalLink xmlns="{NS_MAIN}" xmlns:r="{NS_R}">'
        f'<externalBook r:id="rId1"><sheetNames><sheetName val="{target}"/></sheetNames>'
        '<definedNames><definedName name="ExtOnlyName" refersTo="=Gone!$Z$1"/></definedNames>'
        '<sheetDataSet><sheetData sheetId="0"><row r="1"><cell r="A1"><v>7</v></cell></row>'
        '</sheetData></sheetDataSet></externalBook></externalLink>'
    )
    xml_link_rels = (
        f'{xml_head}<Relationships xmlns="{NS_PKG}">'
        f'<Relationship Id="rId1" Type="{REL}externalLinkPath" '
        'Target="file:///Users/mmf/Source.xlsx" TargetMode="External"/>'
        '</Relationships>'
    )
    xml_content_types = (
        f'{xml_head}<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        '<Override PartName="/xl/worksheets/sheet2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        '<Override PartName="/xl/externalLinks/externalLink1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml"/>'
        '<Override PartName="/xl/calcChain.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/>'
        '</Types>'
    )
    xml_root_rels = (
        f'{xml_head}<Relationships xmlns="{NS_PKG}">'
        f'<Relationship Id="rId1" Type="{REL}officeDocument" Target="xl/workbook.xml"/>'
        '</Relationships>'
    )
    parts = {
        '[Content_Types].xml': xml_content_types,
        '_rels/.rels': xml_root_rels,
        'xl/workbook.xml': xml_workbook,
        'xl/_rels/workbook.xml.rels': xml_workbook_rels,
        'xl/worksheets/sheet1.xml': xml_data,
        'xl/worksheets/sheet2.xml': xml_ext,
        'xl/externalLinks/externalLink1.xml': xml_link,
        'xl/externalLinks/_rels/externalLink1.xml.rels': xml_link_rels,
        'xl/calcChain.xml': '<calcChain><c r="B1" i="1"/></calcChain>',
    }
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, xml in parts.items():
            archive.writestr(name, xml)


def make_processor(config_extra=None) -> SeverExternalTiesProcessor:
    config = {'processor_type': 'sever_external_ties',
              'step_description': 'antagonist',
              'files': ['unused.xlsx']}
    config.update(config_extra or {})
    return SeverExternalTiesProcessor(config)


def sha(path: str) -> str:
    with open(path, 'rb') as handle:
        return hashlib.sha256(handle.read()).hexdigest()


def only_output(workdir: str, stem: str) -> str:
    outputs = [name for name in os.listdir(workdir)
               if name.startswith(stem + '_severed_') and name.endswith('.xlsx')]
    return os.path.join(workdir, outputs[0]) if len(outputs) == 1 else ''


def part_text(path: str, part: str) -> str:
    with zipfile.ZipFile(path) as archive:
        return archive.read(part).decode('utf-8')


def check_all(checks: list) -> bool:
    passed = True
    for ok, label in checks:
        print(f"  {'✓' if ok else '✗'} {label}")
        passed &= bool(ok)
    return passed


def test_hostile_refuses_all_or_nothing():
    """A ref with no local sheet and no cache -> refuse, no output, report."""
    print("\nTesting all-or-nothing refusal on the audit drill file...")
    with tempfile.TemporaryDirectory() as workdir:
        path = os.path.join(workdir, 'hostile.xlsx')
        build_hostile_workbook(path)
        before = sha(path)
        report = os.path.join(workdir, 'sever.json')
        raised = ''
        try:
            make_processor({'files': [path], 'report_file': report}).perform_file_operation()
        except StepProcessorError as error:
            raised = str(error)
        with open(report, encoding='utf-8') as handle:
            loaded = json.load(handle)
        reasons = {entry['where']: entry['reason'] for entry in loaded['refused']}
        return check_all([
            ('no output written' in raised, f"raised: {raised[:70]}"),
            (only_output(workdir, 'hostile') == '', 'no output file'),
            (sha(path) == before, 'source untouched'),
            (reasons.get('Data!D1', '').startswith('no cached value'),
             'D1 refused for missing cache'),
            (set(reasons) == {'Data!D1'}, f"only D1 refused: {sorted(reasons)}"),
            (any(entry['where'] == 'Data!F1' for entry in loaded['fixed']['frozen']),
             'F1 (orphan [3], cached) would have frozen'),
            (loaded['output'] == '' and loaded['plumbing_removed'] == {},
             'report shows nothing removed'),
            (any(entry['kind'] == 'file_hyperlinks' for entry in loaded['limbo']),
             'file hyperlink reported as limbo'),
        ])


def test_copied_tab_retargets_everything():
    """Every [1]Ext ref becomes a local Ext ref; plumbing gone; loads."""
    print("\nTesting the copied-tab retarget path...")
    with tempfile.TemporaryDirectory() as workdir:
        path = os.path.join(workdir, 'copied.xlsx')
        build_linked_workbook(path, local_sheet_for_refs=True)
        result = make_processor({'files': [path]}).perform_file_operation()
        output = only_output(workdir, 'copied')
        data = part_text(output, 'xl/worksheets/sheet1.xml')
        workbook = part_text(output, 'xl/workbook.xml')
        rels = part_text(output, 'xl/_rels/workbook.xml.rels')
        types = part_text(output, '[Content_Types].xml')
        with zipfile.ZipFile(output) as archive:
            names = archive.namelist()
        loaded = openpyxl.load_workbook(output)
        sheet = loaded['Data']
        after = inventory_external_ties(output)
        return check_all([
            ('severed 1 file' in result, f"result: {result}"),
            ('<f>Ext!A1*2</f><v>14</v>' in data, 'bare ref retargeted, cache kept'),
            ("<f>'Ext'!A1&amp;\"x\"</f>" in data, 'quoted ref retargeted'),
            ('<f>LocalName</f>' in data, '[1]!LocalName -> local name'),
            ('<f>[1]!ExtOnlyName</f>' not in data and '<c r="D1"><v>5</v></c>' in data,
             'external-only name frozen (default policy)'),
            ('si="0">Ext!A1+A1</f>' in data and '<f t="shared" si="0"/>' in data,
             'shared master retargeted, slaves untouched'),
            ('<definedName name="ExtName" hidden="1">Ext!$B$2</definedName>' in workbook,
             'defined name retargeted'),
            ('<formula>A1&gt;Ext!A1</formula>' in data, 'CF rule retargeted'),
            ('<formula1>Ext!$A$1:$A$5</formula1>' in data and 'count="2"' in data,
             'DV rule retargeted, count intact'),
            ('externalReferences' not in workbook, 'externalReferences block gone'),
            ('externalLink' not in rels and 'calcChain' not in rels, 'rels gone'),
            ('externalLinks' not in types and 'calcChain' not in types, 'overrides gone'),
            (not any(name.startswith('xl/externalLinks/') or name == 'xl/calcChain.xml'
                     for name in names), 'parts gone'),
            ('fullCalcOnLoad="1"' in workbook and 'calcMode="manual"' not in workbook,
             'recalc forced, manual normalized'),
            (sheet['B1'].value == '=Ext!A1*2' and sheet['F3'].value == '=Ext!A3+A3',
             'openpyxl reads retargeted formulas incl. shared slave'),
            (not after['summary']['has_ties'], 'post-write inventory clean'),
        ])


def test_freeze_and_drop_policies():
    """Nothing retargets: formulas frozen (shared group too), rules dropped."""
    print("\nTesting freeze / drop defaults when nothing retargets...")
    with tempfile.TemporaryDirectory() as workdir:
        path = os.path.join(workdir, 'gone.xlsx')
        build_linked_workbook(path, local_sheet_for_refs=False)
        report = os.path.join(workdir, 'sever.json')
        make_processor({'files': [path], 'report_file': report}).perform_file_operation()
        output = only_output(workdir, 'gone')
        data = part_text(output, 'xl/worksheets/sheet1.xml')
        workbook = part_text(output, 'xl/workbook.xml')
        with open(report, encoding='utf-8') as handle:
            loaded = json.load(handle)
        frozen = {entry['where'] for entry in loaded['fixed']['frozen']}
        dropped = {(entry['kind'], entry['where']) for entry in loaded['fixed']['dropped']}
        sheet = openpyxl.load_workbook(output)['Data']
        return check_all([
            ('<c r="B1"><v>14</v></c>' in data, 'B1 frozen to 14'),
            ('<c r="C1" t="str"><v>7x</v></c>' in data, 'string result keeps t="str"'),
            ('<c r="F1"><v>10</v></c>' in data and '<c r="F2"><v>11</v></c>' in data
             and '<c r="F3"><v>12</v></c>' in data, 'shared master and slaves frozen'),
            ('<f>A1*10</f>' in data and '<f>LocalName</f>' in data,
             'local formulas untouched'),
            (frozen == {'Data!B1', 'Data!C1', 'Data!D1', 'Data!F1', 'Data!F2', 'Data!F3'},
             f"frozen set {sorted(frozen)}"),
            ('ExtName' not in workbook and 'LocalName' in workbook, 'external name dropped'),
            ('sqref="A1:A9"' not in data and 'sqref="G1:G9"' in data,
             'emptied CF wrapper removed, local CF kept'),
            ('count="1"' in data and 'sqref="B2:B9"' not in data and 'sqref="C2:C9"' in data,
             'DV rule dropped and count corrected'),
            (('conditional_formatting', 'Data!A1:A9') in dropped
             and ('data_validation', 'Data!B2:B9') in dropped
             and ('defined_name', 'ExtName') in dropped, 'drops reported'),
            (sheet['B1'].value == 14 and sheet['C1'].value == '7x', 'openpyxl reads values'),
            (loaded['before_summary']['has_ties'] and not loaded['after_summary']['has_ties'],
             'before/after summaries'),
        ])


def test_refuse_policies_and_arrays():
    """Refuse policy blocks output; an array formula is a forced refusal."""
    print("\nTesting refuse policies and array refusal...")
    with tempfile.TemporaryDirectory() as workdir:
        path = os.path.join(workdir, 'gone.xlsx')
        build_linked_workbook(path, local_sheet_for_refs=False)
        raised = ''
        try:
            make_processor({'files': [path],
                            'unresolved_formula_policy': 'refuse'}).perform_file_operation()
        except StepProcessorError as error:
            raised = str(error)
        policy_blocked = 'no output written' in raised and only_output(workdir, 'gone') == ''

        array_path = os.path.join(workdir, 'array.xlsx')
        build_linked_workbook(array_path, local_sheet_for_refs=False)
        xml = part_text(array_path, 'xl/worksheets/sheet1.xml')
        xml = xml.replace('<c r="B1"><f>', '<c r="B1"><f t="array" ref="B1">', 1)
        with zipfile.ZipFile(array_path) as source:
            parts = {name: source.read(name) for name in source.namelist()}
        parts['xl/worksheets/sheet1.xml'] = xml.encode()
        with zipfile.ZipFile(array_path, 'w', zipfile.ZIP_DEFLATED) as target:
            for name, payload in parts.items():
                target.writestr(name, payload)
        raised = ''
        try:
            make_processor({'files': [array_path]}).perform_file_operation()
        except StepProcessorError as error:
            raised = str(error)
        array_blocked = 'no output written' in raised and only_output(workdir, 'array') == ''

        try:
            make_processor({'files': [path], 'unresolved_defined_name_policy': 'freeze'})
            vocab_refused = False
        except StepProcessorError as error:
            vocab_refused = 'unresolved_defined_name_policy' in str(error)
        return check_all([
            (policy_blocked, 'formula refuse policy blocks output'),
            (array_blocked, 'array formula forces refusal'),
            (vocab_refused, 'freeze is not a defined-name policy'),
        ])


def test_output_naming_and_skip():
    """Suffix + timestamp naming, output_dir honored, tie-free file skipped."""
    print("\nTesting output naming and the no-ties skip...")
    with tempfile.TemporaryDirectory() as workdir:
        path = os.path.join(workdir, 'copied.xlsx')
        build_linked_workbook(path, local_sheet_for_refs=True)
        out_dir = os.path.join(workdir, 'out')
        os.mkdir(out_dir)
        make_processor({'files': [path], 'output_dir': out_dir,
                        'output_suffix': '_clean',
                        'timestamp_format': '%Y%m%d'}).perform_file_operation()
        outputs = os.listdir(out_dir)
        pattern_ok = len(outputs) == 1 and re.fullmatch(r'copied_clean_\d{8}\.xlsx', outputs[0])
        try:
            make_processor({'files': [path], 'output_dir': out_dir,
                            'output_suffix': '_clean',
                            'timestamp_format': '%Y%m%d'}).perform_file_operation()
            collision_refused = False
        except StepProcessorError as error:
            collision_refused = 'already exists' in str(error)

        clean = os.path.join(workdir, 'clean.xlsx')
        workbook = openpyxl.Workbook()
        workbook.active['A1'] = '=1+1'
        workbook.save(clean)
        result = make_processor({'files': [clean]}).perform_file_operation()
        return check_all([
            (pattern_ok, f"outputs {outputs}"),
            (collision_refused, 'same-day collision refused rather than overwritten'),
            ('severed 1 file' in result and only_output(workdir, 'clean') == '',
             'tie-free file skipped without an output'),
        ])


def main():
    """Run every test and report a final score."""
    print("=== sever_external_ties antagonistic tests ===")

    tests = [
        test_hostile_refuses_all_or_nothing,
        test_copied_tab_retargets_everything,
        test_freeze_and_drop_policies,
        test_refuse_policies_and_arrays,
        test_output_naming_and_skip,
    ]

    passed = 0
    for test_func in tests:
        try:
            if test_func():
                passed += 1
        except Exception as error:
            print(f"✗ {test_func.__name__} crashed: {error}")

    print(f"\n=== Results: {passed}/{len(tests)} tests passed ===")
    return passed == len(tests)


if __name__ == '__main__':
    sys.exit(0 if main() else 1)

# End of file #
