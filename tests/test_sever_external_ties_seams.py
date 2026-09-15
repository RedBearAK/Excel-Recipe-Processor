"""
Antagonistic tests for sever_external_ties, round two: the seams.

tests/test_sever_external_ties_seams.py

Every case here is one the first module did not cover: two links
where only one retargets, a mixed formula, [N] inside a string
literal, sheet names that need quoting, 3D refs, cached error
values, Excel-order cell attributes, a slave with no cache, an x14
extLst carrier, a chart that retargets and one that does not, an
external pivot source, a workbook with no calcPr, retargeting turned
off, and a multi-file step where one file refuses. Runnable directly
or with pytest.
"""

import os
import sys
import json
import tempfile
import zipfile

import openpyxl

from excel_recipe_processor.core.base_processor import StepProcessorError
from excel_recipe_processor.processors._helpers.external_ties_inventory import (
    inventory_external_ties,
)

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from test_audit_external_ties import (  # noqa: E402
    NS_R,
    REL,
    NS_PKG,
    NS_MAIN,
    build_hostile_workbook,
)
from test_sever_external_ties import (  # noqa: E402
    sha,
    part_text,
    check_all,
    only_output,
    make_processor,
    build_linked_workbook,
)

XML_HEAD = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
CT = 'http://schemas.openxmlformats.org/package/2006/content-types'


def write_package(path: str, parts: dict) -> None:
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, xml in parts.items():
            archive.writestr(name, xml)


def base_parts(sheets: list, data_xml: str, extra_names: str = '',
               calc_pr: str = '<calcPr calcId="1"/>', links: int = 2,
               extra_parts=None, extra_workbook_rels: str = '',
               extra_overrides: str = '') -> dict:
    """A package with N external links and a Data sheet plus named sheets."""
    sheet_elements = ''.join(
        f'<sheet name="{name}" sheetId="{i}" r:id="rId{i}"/>'
        for i, name in enumerate(sheets, start=1))
    sheet_rels = ''.join(
        f'<Relationship Id="rId{i}" Type="{REL}worksheet" Target="worksheets/sheet{i}.xml"/>'
        for i, _ in enumerate(sheets, start=1))
    link_refs = ''.join(
        f'<externalReference r:id="rIdL{n}"/>' for n in range(1, links + 1))
    link_rels = ''.join(
        f'<Relationship Id="rIdL{n}" Type="{REL}externalLink" '
        f'Target="externalLinks/externalLink{n}.xml"/>' for n in range(1, links + 1))
    parts = {
        '[Content_Types].xml': (
            f'{XML_HEAD}<Types xmlns="{CT}">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            + ''.join(f'<Override PartName="/xl/worksheets/sheet{i}.xml" '
                      'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
                      for i, _ in enumerate(sheets, start=1))
            + ''.join(f'<Override PartName="/xl/externalLinks/externalLink{n}.xml" '
                      'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml"/>'
                      for n in range(1, links + 1))
            + extra_overrides + '</Types>'),
        '_rels/.rels': (
            f'{XML_HEAD}<Relationships xmlns="{NS_PKG}">'
            f'<Relationship Id="rId1" Type="{REL}officeDocument" Target="xl/workbook.xml"/>'
            '</Relationships>'),
        'xl/workbook.xml': (
            f'{XML_HEAD}<workbook xmlns="{NS_MAIN}" xmlns:r="{NS_R}">'
            f'<sheets>{sheet_elements}</sheets>'
            f'<externalReferences>{link_refs}</externalReferences>'
            f'<definedNames><definedName name="LocalName">Data!$A$1</definedName>{extra_names}</definedNames>'
            f'{calc_pr}</workbook>'),
        'xl/_rels/workbook.xml.rels': (
            f'{XML_HEAD}<Relationships xmlns="{NS_PKG}">{sheet_rels}{link_rels}'
            f'{extra_workbook_rels}</Relationships>'),
        'xl/worksheets/sheet1.xml': data_xml,
    }
    for i, name in enumerate(sheets[1:], start=2):
        parts[f'xl/worksheets/sheet{i}.xml'] = (
            f'{XML_HEAD}<worksheet xmlns="{NS_MAIN}"><sheetData>'
            f'<row r="1"><c r="A1"><v>{i}</v></c></row></sheetData></worksheet>')
    for n in range(1, links + 1):
        parts[f'xl/externalLinks/externalLink{n}.xml'] = (
            f'{XML_HEAD}<externalLink xmlns="{NS_MAIN}" xmlns:r="{NS_R}">'
            f'<externalBook r:id="rId1"><sheetNames><sheetName val="Link{n}"/>'
            + ''.join(f'<sheetName val="{name}"/>' for name in sheets[1:])
            + '</sheetNames>'
            '</externalBook></externalLink>')
        parts[f'xl/externalLinks/_rels/externalLink{n}.xml.rels'] = (
            f'{XML_HEAD}<Relationships xmlns="{NS_PKG}">'
            f'<Relationship Id="rId1" Type="{REL}externalLinkPath" '
            f'Target="file:///tmp/Source{n}.xlsx" TargetMode="External"/></Relationships>')
    parts.update(extra_parts or {})
    return parts


def data_sheet(cells: str, tail: str = '') -> str:
    return (f'{XML_HEAD}<worksheet xmlns="{NS_MAIN}" xmlns:r="{NS_R}">'
            f'<sheetData><row r="1">{cells}</row></sheetData>{tail}</worksheet>')


def run_default(path: str, extra=None):
    """Run the processor; return (raised_message, output_path)."""
    try:
        make_processor({'files': [path], **(extra or {})}).perform_file_operation()
        raised = ''
    except StepProcessorError as error:
        raised = str(error)
    stem = os.path.splitext(os.path.basename(path))[0]
    return raised, only_output(os.path.dirname(path), stem)


def test_two_links_one_retargets():
    """[1] names a local sheet, [2] does not; cached cells freeze; plumbing all gone."""
    print("\nTesting two links where only one retargets...")
    cells = (
        '<c r="A1"><v>1</v></c>'
        '<c r="B1"><f>[1]Local!A1</f><v>2</v></c>'
        '<c r="C1"><f>[2]Nope!A1</f><v>3</v></c>'
        '<c r="D1"><f>[1]Local!A1+[2]Nope!A1</f><v>5</v></c>'
        '<c r="E1"><f>[12]Local!A1</f><v>9</v></c>'
        '<c r="F1"><f>[1]Data!A1</f><v>1</v></c>'
    )
    with tempfile.TemporaryDirectory() as workdir:
        path = os.path.join(workdir, 'two.xlsx')
        write_package(path, base_parts(['Data', 'Local'], data_sheet(cells)))
        report = os.path.join(workdir, 'r.json')
        raised, output = run_default(path, {'report_file': report})
        data = part_text(output, 'xl/worksheets/sheet1.xml') if output else ''
        with open(report, encoding='utf-8') as handle:
            loaded = json.load(handle)
        frozen = {e['where'] for e in loaded['fixed']['frozen']}
        return check_all([
            (raised == '' and output, 'output written'),
            ('<f>Local!A1</f><v>2</v>' in data, '[1] retargeted'),
            ('<c r="C1"><v>3</v></c>' in data, '[2] frozen'),
            ('<c r="D1"><v>5</v></c>' in data, 'mixed formula frozen, not half-rewritten'),
            ('<c r="E1"><v>9</v></c>' in data, 'orphan two-digit [12] frozen, not retargeted'),
            ('<c r="F1"><v>1</v></c>' in data,
             'local sheet the source never declared is frozen, not retargeted'),
            (frozen == {'Data!C1', 'Data!D1', 'Data!E1', 'Data!F1'}, f"frozen {sorted(frozen)}"),
            (len(loaded['plumbing_removed']['parts']) == 4, 'both link parts and rels gone'),
            (not inventory_external_ties(output)['summary']['has_ties'], 'clean after'),
        ])


def test_string_literal_and_quoting():
    """[1] inside a string literal freezes; names needing quotes stay quoted."""
    print("\nTesting string literals and quoted sheet names...")
    cells = (
        '<c r="A1" t="str"><f>"[1]"&amp;A2</f><v>[1]x</v></c>'
        "<c r=\"B1\"><f>'[1]Van Detail'!A1</f><v>2</v></c>"
        "<c r=\"C1\"><f>'[1]Bob''s'!A1</f><v>3</v></c>"
        "<c r=\"D1\"><f>SUM('[1]Van Detail:Bob''s'!A1)</f><v>5</v></c>"
        "<c r=\"E1\"><f>SUM('[1]Van Detail:Missing'!A1)</f><v>6</v></c>"
    )
    with tempfile.TemporaryDirectory() as workdir:
        path = os.path.join(workdir, 'quote.xlsx')
        write_package(path, base_parts(['Data', 'Van Detail', "Bob's"], data_sheet(cells)))
        raised, output = run_default(path)
        data = part_text(output, 'xl/worksheets/sheet1.xml') if output else ''
        loaded = openpyxl.load_workbook(output)['Data'] if output else None
        return check_all([
            (raised == '' and output, 'output written'),
            ('<c r="A1" t="str"><v>[1]x</v></c>' in data,
             'string-literal [1] frozen (cannot be proven not a reference)'),
            ("<f>'Van Detail'!A1</f>" in data, 'space-named sheet keeps quotes'),
            ("<f>'Bob''s'!A1</f>" in data, 'escaped apostrophe survives'),
            ("<f>SUM('Van Detail:Bob''s'!A1)</f>" in data, '3D ref with both ends local'),
            ('<c r="E1"><v>6</v></c>' in data, '3D ref with a missing end frozen'),
            (loaded is not None and loaded['B1'].value == "='Van Detail'!A1",
             'openpyxl reads the quoted formula'),
        ])


def test_excel_attribute_order_and_error_cache():
    """Attributes in Excel order, cached #N/A, and inline-string cell."""
    print("\nTesting attribute order, error caches, inline strings...")
    cells = (
        '<c t="e" r="A1" s="3"><f>[1]Gone!A1</f><v>#N/A</v></c>'
        '<c s="2" r="B1"><f>[1]Local!A1</f><v>4</v></c>'
        '<c r="C1" t="inlineStr"><is><t>plain</t></is></c>'
        '<c r="D1" s="1"/>'
        '<c r="E1"><f>[1]Gone!A1</f><v></v></c>'
    )
    with tempfile.TemporaryDirectory() as workdir:
        path = os.path.join(workdir, 'order.xlsx')
        write_package(path, base_parts(['Data', 'Local'], data_sheet(cells), links=1))
        raised, output = run_default(path)
        data = part_text(output, 'xl/worksheets/sheet1.xml') if output else ''
        return check_all([
            (raised == '' and output, 'output written'),
            ('<c t="e" r="A1" s="3"><v>#N/A</v></c>' in data,
             'error cache frozen with t="e" and style kept'),
            ('<c s="2" r="B1"><f>Local!A1</f><v>4</v></c>' in data,
             'retarget with attributes in Excel order'),
            ('<c r="C1" t="inlineStr"><is><t>plain</t></is></c>' in data,
             'inline string untouched'),
            ('<c r="D1" s="1"/>' in data, 'self-closing styled cell untouched'),
            ('<c r="E1"><v></v></c>' in data, 'empty <v> counts as a cache (Excel wrote it)'),
        ])


def test_slave_without_cache_refuses_group():
    """A shared slave with no <v> refuses; the master is reported too."""
    print("\nTesting shared slave without a cached value...")
    cells = (
        '<c r="A1"><f t="shared" ref="A1:A3" si="0">[1]Gone!A1</f><v>1</v></c>'
        '<c r="A2"><f t="shared" si="0"/><v>2</v></c>'
        '<c r="A3"><f t="shared" si="0"/></c>'
    )
    with tempfile.TemporaryDirectory() as workdir:
        path = os.path.join(workdir, 'slave.xlsx')
        write_package(path, base_parts(['Data'], data_sheet(cells), links=1))
        before = sha(path)
        report = os.path.join(workdir, 'r.json')
        raised, output = run_default(path, {'report_file': report})
        with open(report, encoding='utf-8') as handle:
            loaded = json.load(handle)
        refused = {e['where']: e['reason'] for e in loaded['refused']}
        return check_all([
            ('no output written' in raised and not output, 'refused, no output'),
            (sha(path) == before, 'source untouched'),
            ('Data!A3' in refused and 'shared slave' in refused['Data!A3'],
             'A3 named with its reason'),
            (any(e['where'] == 'Data!A1' for e in loaded['fixed']['frozen']),
             'report still shows what A1 would have done'),
        ])


def test_unknown_carrier_refuses():
    """An x14 extLst data validation with [N] is unclassified -> refuse."""
    print("\nTesting unknown carrier refusal (x14 extLst)...")
    tail = (
        '<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}" '
        'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
        '<x14:dataValidations count="1"><x14:dataValidation type="list" sqref="A2:A9">'
        '<x14:formula1><xm:f xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
        '[1]Local!$A$1:$A$5</xm:f></x14:formula1></x14:dataValidation>'
        '</x14:dataValidations></ext></extLst>'
    )
    cells = '<c r="A1"><f>[1]Local!A1</f><v>1</v></c>'
    with tempfile.TemporaryDirectory() as workdir:
        path = os.path.join(workdir, 'x14.xlsx')
        write_package(path, base_parts(['Data', 'Local'], data_sheet(cells, tail), links=1))
        report = os.path.join(workdir, 'r.json')
        raised, output = run_default(path, {'report_file': report})
        with open(report, encoding='utf-8') as handle:
            loaded = json.load(handle)
        kinds = {e['kind'] for e in loaded['refused']}
        return check_all([
            ('no output written' in raised and not output, 'refused, no output'),
            (any(k.startswith('other:') or k == 'unclassified' for k in kinds),
             f"refusal kinds {sorted(kinds)}"),
            (any('x14' in e['reason'] for e in loaded['refused']),
             'reason names the carrier element'),
        ])


def test_charts_and_external_pivot():
    """A chart series retargets; one that does not refuses; external pivot refuses."""
    print("\nTesting chart series and external pivot source...")
    chart_ok = (f'{XML_HEAD}<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart">'
                '<c:ser><c:val><c:numRef><c:f>[1]Local!$A$1:$A$3</c:f></c:numRef></c:val></c:ser>'
                '</c:chartSpace>')
    chart_bad = chart_ok.replace('[1]Local', '[1]Gone')
    cells = '<c r="A1"><f>[1]Local!A1</f><v>1</v></c>'
    with tempfile.TemporaryDirectory() as workdir:
        path = os.path.join(workdir, 'chart.xlsx')
        write_package(path, base_parts(['Data', 'Local'], data_sheet(cells), links=1,
                                       extra_parts={'xl/charts/chart1.xml': chart_ok}))
        raised_ok, output_ok = run_default(path)
        chart_after = part_text(output_ok, 'xl/charts/chart1.xml') if output_ok else ''

        path2 = os.path.join(workdir, 'chartbad.xlsx')
        write_package(path2, base_parts(['Data', 'Local'], data_sheet(cells), links=1,
                                        extra_parts={'xl/charts/chart1.xml': chart_bad}))
        raised_bad, output_bad = run_default(path2)

        pivot_parts = {
            'xl/pivotCache/pivotCacheDefinition1.xml': (
                f'{XML_HEAD}<pivotCacheDefinition xmlns="{NS_MAIN}" xmlns:r="{NS_R}">'
                '<cacheSource type="worksheet"><worksheetSource ref="A1:B9" sheet="Link1" r:id="rId1"/>'
                '</cacheSource></pivotCacheDefinition>'),
            'xl/pivotCache/_rels/pivotCacheDefinition1.xml.rels': (
                f'{XML_HEAD}<Relationships xmlns="{NS_PKG}">'
                f'<Relationship Id="rId1" Type="{REL}externalLink" '
                'Target="../externalLinks/externalLink1.xml"/></Relationships>'),
            'xl/pivotTables/pivotTable1.xml': (
                f'{XML_HEAD}<pivotTableDefinition xmlns="{NS_MAIN}" name="P">'
                '<location ref="C1:D5"/></pivotTableDefinition>'),
            'xl/pivotTables/_rels/pivotTable1.xml.rels': (
                f'{XML_HEAD}<Relationships xmlns="{NS_PKG}">'
                f'<Relationship Id="rId1" Type="{REL}pivotCacheDefinition" '
                'Target="../pivotCache/pivotCacheDefinition1.xml"/></Relationships>'),
            'xl/worksheets/_rels/sheet1.xml.rels': (
                f'{XML_HEAD}<Relationships xmlns="{NS_PKG}">'
                f'<Relationship Id="rId1" Type="{REL}pivotTable" '
                'Target="../pivotTables/pivotTable1.xml"/></Relationships>'),
        }
        path3 = os.path.join(workdir, 'pivot.xlsx')
        write_package(path3, base_parts(['Data', 'Local'], data_sheet(cells), links=1,
                                        extra_parts=pivot_parts))
        report = os.path.join(workdir, 'r.json')
        raised_pivot, output_pivot = run_default(path3, {'report_file': report})
        with open(report, encoding='utf-8') as handle:
            loaded = json.load(handle)
        return check_all([
            (raised_ok == '' and '<c:f>Local!$A$1:$A$3</c:f>' in chart_after,
             'chart series retargeted'),
            ('no output written' in raised_bad and not output_bad,
             'non-retargeting chart refuses'),
            ('no output written' in raised_pivot and not output_pivot,
             'external pivot source refuses'),
            (any(e['kind'] == 'pivot_source' for e in loaded['refused']),
             'pivot refusal in report'),
        ])


def test_no_calc_pr_and_retarget_off():
    """No calcPr -> one is added; retarget_if_possible false freezes all."""
    print("\nTesting missing calcPr and retargeting disabled...")
    cells = '<c r="A1"><f>[1]Local!A1</f><v>1</v></c>'
    with tempfile.TemporaryDirectory() as workdir:
        path = os.path.join(workdir, 'nocalc.xlsx')
        write_package(path, base_parts(['Data', 'Local'], data_sheet(cells), calc_pr='', links=1))
        raised, output = run_default(path)
        workbook = part_text(output, 'xl/workbook.xml') if output else ''

        path2 = os.path.join(workdir, 'off.xlsx')
        write_package(path2, base_parts(['Data', 'Local'], data_sheet(cells), links=1))
        raised2, output2 = run_default(path2, {'retarget_if_possible': False})
        data2 = part_text(output2, 'xl/worksheets/sheet1.xml') if output2 else ''
        return check_all([
            (raised == '' and '<calcPr fullCalcOnLoad="1" calcCompleted="0"/>' in workbook,
             'calcPr added when absent'),
            (workbook.count('<calcPr') == 1, 'exactly one calcPr'),
            (raised2 == '' and '<c r="A1"><v>1</v></c>' in data2,
             'retargetable cell frozen when retargeting is off'),
        ])


def test_multi_file_one_refuses():
    """Good file is written, bad file is not, step still fails, report is a list."""
    print("\nTesting a multi-file step with one refusal...")
    with tempfile.TemporaryDirectory() as workdir:
        good = os.path.join(workdir, 'good.xlsx')
        bad = os.path.join(workdir, 'bad.xlsx')
        build_linked_workbook(good, local_sheet_for_refs=True)
        build_hostile_workbook(bad)
        report = os.path.join(workdir, 'r.json')
        try:
            make_processor({'files': [good, bad],
                            'report_file': report}).perform_file_operation()
            raised = ''
        except StepProcessorError as error:
            raised = str(error)
        with open(report, encoding='utf-8') as handle:
            loaded = json.load(handle)
        return check_all([
            ('bad.xlsx' in raised and 'good.xlsx' not in raised, 'raised names bad only'),
            (only_output(workdir, 'good') != '', 'good output written'),
            (only_output(workdir, 'bad') == '', 'bad output absent'),
            (isinstance(loaded, list) and len(loaded) == 2, 'report is a two-entry list'),
            (loaded[0]['output'] and not loaded[1]['output'], 'per-file output fields'),
        ])


def test_openpyxl_round_trip_of_output():
    """openpyxl can load AND re-save the severed file without complaint."""
    print("\nTesting openpyxl round trip of a severed file...")
    with tempfile.TemporaryDirectory() as workdir:
        path = os.path.join(workdir, 'gone.xlsx')
        build_linked_workbook(path, local_sheet_for_refs=False)
        raised, output = run_default(path)
        resaved = os.path.join(workdir, 'resaved.xlsx')
        error = ''
        try:
            workbook = openpyxl.load_workbook(output)
            workbook.save(resaved)
            again = openpyxl.load_workbook(resaved)
            value = again['Data']['B1'].value
        except Exception as caught:
            error = str(caught)
            value = None
        return check_all([
            (raised == '' and output, 'severed output exists'),
            (error == '', f"round trip error: {error or 'none'}"),
            (value == 14, 'frozen value survives the round trip'),
        ])


def test_link_without_sheet_list_and_entities():
    """No <sheetNames> in the link part -> retarget allowed; entities survive."""
    print("\nTesting a link part with no sheet list, plus XML entities...")
    cells = (
        '<c r="A1"><f>IF([1]Local!A1&lt;3,"a&amp;b",[1]Local!B1)</f><v>7</v></c>'
        '<c r="B1"><f>[1]Local!A1</f><v>1</v></c>'
    )
    with tempfile.TemporaryDirectory() as workdir:
        path = os.path.join(workdir, 'nolist.xlsx')
        parts = base_parts(['Data', 'Local'], data_sheet(cells), links=1)
        parts['xl/externalLinks/externalLink1.xml'] = (
            f'{XML_HEAD}<externalLink xmlns="{NS_MAIN}" xmlns:r="{NS_R}">'
            '<externalBook r:id="rId1"/></externalLink>')
        write_package(path, parts)
        raised, output = run_default(path)
        data = part_text(output, 'xl/worksheets/sheet1.xml') if output else ''
        return check_all([
            (raised == '' and output, 'output written'),
            ('<f>IF(Local!A1&lt;3,"a&amp;b",Local!B1)</f>' in data,
             'two refs retargeted inside one formula, entities untouched'),
            ('<f>Local!A1</f><v>1</v>' in data, 'plain ref retargeted with no sheet list'),
        ])


def main():
    """Run every test and report a final score."""
    print("=== sever_external_ties seam tests ===")

    tests = [
        test_two_links_one_retargets,
        test_string_literal_and_quoting,
        test_excel_attribute_order_and_error_cache,
        test_slave_without_cache_refuses_group,
        test_unknown_carrier_refuses,
        test_charts_and_external_pivot,
        test_no_calc_pr_and_retarget_off,
        test_multi_file_one_refuses,
        test_openpyxl_round_trip_of_output,
        test_link_without_sheet_list_and_entities,
    ]

    passed = 0
    for test_func in tests:
        try:
            if test_func():
                passed += 1
        except Exception as error:
            print(f"✗ {test_func.__name__} crashed: {error}")

    print(f"\n=== Results: {passed}/{len(tests)} tests passed ===")
    return passed == len(tests)


if __name__ == '__main__':
    sys.exit(0 if main() else 1)

# End of file #
